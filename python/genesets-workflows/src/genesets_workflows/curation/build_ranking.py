"""Build a curated ranking file from a declarative recipe.

A recipe (`curation/rankings/<NAME>.recipe.yaml`) says where a source table is,
which columns hold the gene and the score, and how to turn the source's ids
and scores into the corpus format: HGNC symbols, one score per gene, written
by `ranking.write_ranking`. Every download is pinned by SHA-256, including the
HGNC release used for id mapping, so a recipe rebuilds the same file.

Recipe keys::

    name: SCHIZOPHRENIA_SCHEMA_EXOME      # output is <name>.rnk.gz beside the recipe
    source:
      url: https://...                     # the table, or an archive holding it
      sha256: <hex>
      member: path/in/archive.txt         # optional, for .zip archives
    table:
      format: tsv | csv | xlsx | whitespace | atlas_json
      sheet: Sheet1                        # xlsx only
      header_row: 1                        # 1-based; rows above it are skipped
    gene_column: gene_id
    gene_id_type: ensembl | entrez | symbol
    fallback_symbol_column: SYMBOL       # optional: when the id does not map,
                                         # accept this column if it is an
                                         # approved symbol (never prev/alias)
    score_column: P meta
    transform: identity | neg_log10 | negate
    missing: ["", "NA"]                   # score values that mean no measurement
    duplicates: max | min | drop | prefer_exact
                                         # several source rows for one symbol:
                                         # keep the max/min score, discard them
                                         # all, or (prefer_exact) let a row whose
                                         # source name is the approved symbol win
                                         # over rows reached through a previous or
                                         # alias name, keeping the max among equals
    store: false                           # commit the built file? Only when the
                                         # source has no stable URL to rebuild from
    hgnc:
      url: https://.../hgnc_complete_set_2026-10-06.txt
      sha256: <hex>

A GWAS ATLAS (atlas.ctglab.nl, Watanabe et al. 2019) MAGMA table has no
download link; its study page loads the table from a session- and
CSRF-protected endpoint. ``source: {kind: ctg_atlas, atlas_id: 4167, sha256:
...}`` replays that request, and ``table: {format: atlas_json}`` reads the
returned rows with the MAGMA column names.
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import http.cookiejar
import io
import json
import math
import re
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

from genesets_workflows.curation.ranking import write_ranking

ATLAS = "https://atlas.ctglab.nl/traitDB"
ATLAS_MAGMA_COLUMNS = ["GENE", "CHR", "START", "STOP", "NSNPS", "NPARAM", "N", "ZSTAT", "P", "SYMBOL"]

TRANSFORMS = {
    "identity": lambda value: value,
    "neg_log10": lambda value: -math.log10(value),
    "negate": lambda value: -value,
}


@dataclass
class Recipe:
    name: str
    source: dict[str, Any]
    table: dict[str, Any]
    gene_column: str
    gene_id_type: str
    score_column: str
    hgnc: dict[str, Any]
    transform: str = "identity"
    missing: list[str] = field(default_factory=lambda: ["", "NA", "NaN", "nan"])
    duplicates: str = "max"
    fallback_symbol_column: str | None = None
    store: bool = False

    @classmethod
    def load(cls, path: Path) -> Recipe:
        return cls(**yaml.safe_load(Path(path).read_text()))


@dataclass
class BuildReport:
    rows: int = 0
    missing: int = 0
    fallback: int = 0
    unmapped: int = 0
    collapsed: int = 0
    ranked: int = 0
    sha256: str = ""

    def summary(self) -> str:
        return (
            f"{self.rows} source rows: {self.missing} without a score, {self.fallback} "
            f"mapped by their approved symbol after their id failed, {self.unmapped} "
            f"without an approved HGNC symbol, {self.collapsed} duplicate symbols collapsed; "
            f"{self.ranked} genes ranked"
        )


def fetch(url: str, sha256: str, cache_dir: Path) -> bytes:
    """Download ``url`` once into ``cache_dir`` and check its pinned SHA-256."""

    def download() -> bytes:
        request = urllib.request.Request(url, headers={"User-Agent": "genesets-curation/0.1"})
        with urllib.request.urlopen(request, timeout=300) as response:
            return response.read()

    return _cached(url, sha256, cache_dir, download)


def fetch_atlas_magma(atlas_id: int, sha256: str, cache_dir: Path) -> bytes:
    """A GWAS ATLAS study's MAGMA gene table, as the study page requests it."""

    def download() -> bytes:
        opener = urllib.request.build_opener(
            urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar())
        )
        page = opener.open(ATLAS, timeout=120).read().decode()
        token = re.search(r'name="csrf-token" content="([^"]+)"', page)
        if token is None:
            raise ValueError(f"{ATLAS}: no csrf-token on the page")
        form = urllib.parse.urlencode(
            {
                "id": str(atlas_id),
                "infile": "magma.genes.out",
                "header": ":".join(ATLAS_MAGMA_COLUMNS),
            }
        ).encode()
        request = urllib.request.Request(
            f"{ATLAS}/DTfile",
            data=form,
            headers={"X-CSRF-TOKEN": token.group(1), "X-Requested-With": "XMLHttpRequest"},
        )
        return opener.open(request, timeout=300).read()

    return _cached(f"{ATLAS}/DTfile?id={atlas_id}", sha256, cache_dir, download)


def _cached(key: str, sha256: str, cache_dir: Path, download) -> bytes:
    cache = cache_dir / hashlib.sha256(key.encode()).hexdigest()[:16]
    if not cache.exists():
        cache_dir.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(download())
    data = cache.read_bytes()
    actual = hashlib.sha256(data).hexdigest()
    if actual != sha256:
        raise ValueError(f"{key}: sha256 is {actual}, recipe pins {sha256}")
    return data


def _decompress(data: bytes) -> bytes:
    # bgzip files are multi-member gzip, which gzip.decompress reads.
    return gzip.decompress(data) if data[:2] == b"\x1f\x8b" else data


def read_table(data: bytes, table: dict[str, Any]) -> list[dict[str, str]]:
    """Rows of the source table as dicts keyed by header names."""
    fmt = table["format"]
    header_row = int(table.get("header_row", 1))
    if fmt == "atlas_json":
        rows = [ATLAS_MAGMA_COLUMNS, *json.loads(data)["data"]]
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook[table["sheet"]] if "sheet" in table else workbook.active
        rows = [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    else:
        text = _decompress(data).decode("utf-8-sig")
        if fmt == "whitespace":
            rows = [line.split() for line in text.splitlines()]
        else:
            delimiter = {"tsv": "\t", "csv": ","}[fmt]
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    header = [name.strip() for name in rows[header_row - 1]]
    return [
        dict(zip(header, (value.strip() for value in row)))
        for row in rows[header_row:]
        if any(value.strip() for value in row)
    ]


class HgncMapper:
    """Maps Ensembl, Entrez or symbol ids to approved HGNC symbols.

    Symbols resolve to themselves when approved, otherwise through a previous
    symbol, then an alias, but only when that name points to one approved gene.
    """

    def __init__(self, hgnc_tsv: str):
        self.by_id: dict[str, dict[str, str]] = {"ensembl": {}, "entrez": {}, "symbol": {}}
        previous: dict[str, set[str]] = {}
        aliases: dict[str, set[str]] = {}
        for row in csv.DictReader(io.StringIO(hgnc_tsv), delimiter="\t"):
            if row["status"] != "Approved":
                continue
            symbol = row["symbol"]
            self.by_id["symbol"][symbol] = symbol
            if row["ensembl_gene_id"]:
                self.by_id["ensembl"][row["ensembl_gene_id"]] = symbol
            if row["entrez_id"]:
                self.by_id["entrez"][row["entrez_id"]] = symbol
            for column, index in (("prev_symbol", previous), ("alias_symbol", aliases)):
                for name in filter(None, (part.strip() for part in row[column].split("|"))):
                    index.setdefault(name, set()).add(symbol)
        self.previous = {name: next(iter(s)) for name, s in previous.items() if len(s) == 1}
        self.aliases = {name: next(iter(s)) for name, s in aliases.items() if len(s) == 1}

    def approved(self, symbol: str) -> str | None:
        return self.by_id["symbol"].get(symbol)

    def map(self, gene: str, id_type: str) -> str | None:
        if id_type == "ensembl":
            gene = gene.split(".", 1)[0]
        if id_type == "entrez":
            gene = gene.removesuffix(".0")
        found = self.by_id[id_type].get(gene)
        if found is None and id_type == "symbol":
            found = self.previous.get(gene) or self.aliases.get(gene)
        return found


def build(recipe_path: Path, cache_dir: Path) -> BuildReport:
    """Build ``<name>.rnk.gz`` beside the recipe and report what was dropped."""
    recipe = Recipe.load(recipe_path)
    if recipe.transform not in TRANSFORMS:
        raise ValueError(f"transform must be one of {sorted(TRANSFORMS)}")
    if recipe.duplicates not in ("max", "min", "drop", "prefer_exact"):
        raise ValueError("duplicates must be max, min, drop or prefer_exact")
    hgnc = HgncMapper(fetch(recipe.hgnc["url"], recipe.hgnc["sha256"], cache_dir).decode())
    if recipe.source.get("kind") == "ctg_atlas":
        data = fetch_atlas_magma(recipe.source["atlas_id"], recipe.source["sha256"], cache_dir)
    else:
        data = fetch(recipe.source["url"], recipe.source["sha256"], cache_dir)
    if "member" in recipe.source:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            data = archive.read(recipe.source["member"])

    transform = TRANSFORMS[recipe.transform]
    keep = min if recipe.duplicates == "min" else max
    ambiguous: set[str] = set()
    exact: dict[str, bool] = {}
    missing = set(recipe.missing)
    report = BuildReport()
    scores: dict[str, float] = {}
    for row in read_table(data, recipe.table):
        report.rows += 1
        raw = row[recipe.score_column]
        if raw in missing:
            report.missing += 1
            continue
        symbol = hgnc.map(row[recipe.gene_column], recipe.gene_id_type)
        if symbol is None and recipe.fallback_symbol_column:
            symbol = hgnc.approved(row[recipe.fallback_symbol_column])
            report.fallback += symbol is not None
        if symbol is None:
            report.unmapped += 1
            continue
        score = transform(float(raw))
        is_exact = row[recipe.gene_column] == symbol
        if symbol in scores:
            report.collapsed += 1
            if recipe.duplicates == "drop":
                ambiguous.add(symbol)
            elif recipe.duplicates == "prefer_exact" and is_exact != exact[symbol]:
                if not is_exact:
                    continue
                exact[symbol] = True
                scores[symbol] = score
                continue
            score = keep(score, scores[symbol])
        exact[symbol] = exact.get(symbol, False) or is_exact
        scores[symbol] = score
    for symbol in ambiguous:
        del scores[symbol]
    report.ranked = len(scores)
    report.sha256 = write_ranking(scores, recipe_path.parent / f"{recipe.name}.rnk.gz")
    return report
